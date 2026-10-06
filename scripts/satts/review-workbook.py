"""
SATTS dictionary review workbook (upgrade plan P8 item 1).

Regenerates the curriculum reviewer's workbook from the migrated policy pack
(`lib/speech/scientific/policy/ar-v1/*.json`), organised by role and domain,
with the natural, accessible and neutral (literal) readings side by side and
two reviewer columns: «موافق؟» (نعم / لا) and «التصحيح».

Read-only on the pack. Writes one .xlsx (default:
docs/frds/satts-dictionary-review-v2.xlsx). The reviewer's answers are applied
by `apply-approvals.py`, never by hand.

Usage (openpyxl from the zakrly-backend venv):
  ../zakrly-backend/.venv/bin/python scripts/satts/review-workbook.py [out.xlsx]
"""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.join(os.path.dirname(__file__), '..', '..')
PACK = os.path.join(ROOT, 'lib', 'speech', 'scientific', 'policy', 'ar-v1')
DEFAULT_OUT = os.path.join(ROOT, 'docs', 'frds', 'satts-dictionary-review-v2.xlsx')

ROLE_AR = {
    'variable': 'متغير',
    'element': 'عنصر كيميائي',
    'greek': 'حرف يوناني',
    'function': 'دالة',
    'fraction': 'كسر',
    'operator': 'عملية',
    'label': 'كلمة بنائية',
    'symbol': 'اسم رمز',
    'number': 'عدد بنائي',
    'repeat': 'تكرار',
    'unit': 'وحدة',
    'prefix': 'بادئة',
    'state': 'حالة مادة',
    'bond': 'رابطة',
    'reaction': 'رمز تفاعل',
}
DOMAIN_AR = {'MATH': 'رياضيات', 'PHYSICS': 'فيزياء', 'CHEMISTRY': 'كيمياء'}

# Sheet order: (title, predicate on (file, roles, domains)).
SHEETS = [
    ('حروف الرياضيات', lambda f, r, d: f == 'letters.json'),
    ('حروف الفيزياء والكيمياء', lambda f, r, d: f == 'chem-letters.json'),
    ('الحروف العربية', lambda f, r, d: f == 'arabic-letters.json'),
    ('الحروف اليونانية', lambda f, r, d: f == 'greek.json'),
    ('الدوال', lambda f, r, d: f == 'functions.json'),
    ('الكسور', lambda f, r, d: f == 'fractions.json'),
    ('الأعداد', lambda f, r, d: f == 'numbers-structural.json'),
    ('العمليات', lambda f, r, d: f == 'operators.json' and 'operator' in r),
    ('الكلمات البنائية', lambda f, r, d: f == 'operators.json' and 'label' in r and not d),
    ('كلمات الكيمياء', lambda f, r, d: f == 'operators.json' and 'label' in r and d),
    ('الروابط', lambda f, r, d: f == 'operators.json' and 'bond' in r),
    ('رموز التفاعل', lambda f, r, d: f == 'operators.json' and 'reaction' in r),
    ('أسماء الرموز', lambda f, r, d: f == 'operators.json' and 'symbol' in r),
    ('الوحدات', lambda f, r, d: f == 'units.json'),
    ('البادئات', lambda f, r, d: f == 'prefixes.json'),
    ('حالات المادة', lambda f, r, d: f == 'states.json'),
]

# Entries the reviewer must look at closely (decisions and interpretations).
# Keyed by (file, key) or, where one key has several roles, (file, key, role).
FLAGGED = {
    ('letters.json', 'v'): 'ڤ: قد لا ينطقها الصوت؛ «في» محجوزة لعلامة الضرب (O-8).',
    ('letters.json', 'V'): 'ڤ: قد لا ينطقها الصوت؛ «في» محجوزة لعلامة الضرب (O-8).',
    ('chem-letters.json', 'V'): 'ڤ: قد لا ينطقها الصوت؛ «في» محجوزة لعلامة الضرب (O-8).',
    ('chem-letters.json', 'v'): 'ڤ: قد لا ينطقها الصوت؛ «في» محجوزة لعلامة الضرب (O-8).',
    ('operators.json', '·'): 'O-8: «نقطة» (وليس «ضرب قياسي»).',
    ('operators.json', '*'): 'O-8: «في» (وليس «ضرب اتجاهي»).',
    ('operators.json', 'given'): 'الخط العمودي في P(A|B): «بشرط» — قراءة الرمز فقط.',
    ('operators.json', 'whole'): 'O-6: «سين زائد واحد، الكل تربيع».',
    ('operators.json', 'hydrate'): 'نقطة الماء في البلورات: «مع» أو «نقطة»؟',
    ('operators.json', 'binom'): 'رمز التوافيق: «توافيق».',
    ('operators.json', 'over-arrow'): 'شرط فوق السهم غير Δ (عامل حفاز، درجة حرارة): يُقرأ برموزه + «فوق السهم». Δ وحدها «بالتسخين».',
    ('operators.json', '-', 'bond'): 'DEC-052: «رابطة أحادية» — هل يقولها المعلم فعلًا، أم يكتفي بقراءة الذرتين؟',
    ('operators.json', '=', 'bond'): 'DEC-052: «رابطة ثنائية» بين ذرتين في صيغة.',
    ('operators.json', '≡', 'bond'): 'DEC-052: «رابطة ثلاثية» بين ذرتين في صيغة.',
    ('operators.json', '→', 'reaction'): 'DEC-052: سهم التفاعل «ينتج» (داخل معادلة كيميائية فقط).',
    ('operators.json', '⇌', 'reaction'): 'DEC-052: «في حالة اتزان مع».',
    ('operators.json', '↑', 'reaction'): 'DEC-052: «يتصاعد» بعد الناتج الغازي.',
    ('operators.json', '↓', 'reaction'): 'DEC-052: «يترسب» بعد الراسب — اختيار مقترح، أكّده.',
    ('operators.json', 'subscript'): 'DEC-053: «تحت» لدليل حرفي في رمز فيزيائي/كيميائي: «كيه تحت بي في تي». الرياضيات والدليل الرقمي بدون «تحت».',
    ('operators.json', 'Δ', 'reaction'): 'DEC-052: Δ فوق سهم التفاعل «بالتسخين»، وتُقال قبل «ينتج».',
}


def flag(file, roles, key):
    for role in roles:
        if (file, key, role) in FLAGGED:
            return FLAGGED[(file, key, role)]
    return FLAGGED.get((file, key))

HEADERS = [
    'الرمز',
    'الدور',
    'المادة',
    'القراءة الطبيعية',
    'القراءة التفصيلية',
    'اسم الرمز المحايد',
    'ملاحظة للمراجع',
    'موافق؟',
    'التصحيح',
    'الملف',
    'المعرّف',
]


def entry_id(file, roles, domains, key):
    """Stable row identity, read back by apply-approvals.py."""
    return f"{file}|{','.join(roles)}|{','.join(domains or [])}|{key}"


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_OUT
    manifest = json.load(open(os.path.join(PACK, 'manifest.json'), encoding='utf-8'))
    rows = []
    for file in manifest['files']:
        table = json.load(open(os.path.join(PACK, file), encoding='utf-8'))
        for entry in table['entries']:
            roles = entry.get('roles') or table.get('roles') or []
            domains = entry.get('domains') or table.get('domains') or []
            rows.append((file, roles, domains, entry))

    wb = Workbook()
    readme = wb.active
    readme.title = 'اقرأني'
    readme.sheet_view.rightToLeft = True
    lines = [
        f"مراجعة قاموس النطق العلمي — {manifest['policyVersion']} (الحالة: {manifest['status']})",
        '',
        'القاعدة: SATTS يغيّر النطق فقط، لا المعنى. لا أسماء مركبات ولا أيونات ولا كميات فيزيائية.',
        'الأعداد بصيغة «اثنين» (O-1). الحرف الكبير والصغير لا يُنطقان بعلامة (O-3).',
        '',
        'لكل سطر: اكتب «نعم» في عمود «موافق؟» إذا كانت القراءة صحيحة.',
        'إذا كانت خاطئة: اكتب «لا» واكتب القراءة الصحيحة في عمود «التصحيح».',
        'لا تغيّر عمودي «الملف» و«المعرّف».',
        '',
        'الأسطر الملوّنة بالأصفر تحتاج انتباهًا خاصًا (قرار أو تفسير).',
        'بعد الانتهاء أرسل الملف؛ تُطبَّق الموافقات بأداة apply-approvals ثم يُعاد قفل القاموس.',
        '',
        f'عدد المدخلات: {len(rows)}',
    ]
    for i, line in enumerate(lines, start=1):
        readme.cell(row=i, column=1, value=line)
    readme.column_dimensions['A'].width = 110

    flag_fill = PatternFill('solid', fgColor='FFF2CC')
    header_fill = PatternFill('solid', fgColor='DDEBF7')
    used = set()
    for title, predicate in SHEETS:
        chosen = [r for r in rows if predicate(r[0], r[1], r[2]) and id(r[3]) not in used]
        if not chosen:
            continue
        ws = wb.create_sheet(title)
        ws.sheet_view.rightToLeft = True
        ws.append(HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        validation = DataValidation(type='list', formula1='"نعم,لا"', allow_blank=True)
        ws.add_data_validation(validation)
        for file, roles, domains, entry in chosen:
            used.add(id(entry))
            note = flag(file, roles, entry['key']) or entry.get('notes', '')
            ws.append([
                entry['key'],
                '، '.join(ROLE_AR.get(r, r) for r in roles),
                '، '.join(DOMAIN_AR.get(d, d) for d in domains) or 'كل المواد',
                entry.get('natural', ''),
                entry.get('accessible', ''),
                entry.get('literal', ''),
                note,
                'نعم' if entry.get('status') == 'approved' else '',
                '',
                file,
                entry_id(file, roles, domains, entry['key']),
            ])
            row = ws.max_row
            validation.add(f'H{row}')
            if flag(file, roles, entry['key']):
                for cell in ws[row]:
                    cell.fill = flag_fill
        for col, width in zip('ABCDEFGHIJK', [10, 16, 16, 26, 30, 22, 48, 9, 26, 22, 36]):
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
    missing = [r for r in rows if id(r[3]) not in used]
    if missing:
        raise SystemExit(f'{len(missing)} entries have no sheet: {[m[3]["key"] for m in missing][:10]}')
    wb.save(out)
    print(f'wrote {out}: {len(rows)} entries, {len(wb.sheetnames) - 1} sheets')


if __name__ == '__main__':
    main()
