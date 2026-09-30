"""
Applies a curriculum reviewer's returned workbook to the policy pack
(upgrade plan P8 item 2). Human gate: only the reviewer's answers approve an
entry, and only with the reviewer's name and date supplied on the command line.

- «موافق؟» = نعم                → status approved (wording unchanged)
- «موافق؟» = لا + «التصحيح»     → natural reading replaced, status approved
- anything else                  → left proposed

Dry run by default: prints every change. `--write` applies them, bumps the
policy version and leaves the manifest `experimental` unless
`--approve-manifest` is given AND no proposed entry remains. Then run
`npx tsx scripts/satts/relock-policy.ts` to relock the content hash.

Usage:
  ../zakrly-backend/.venv/bin/python scripts/satts/apply-approvals.py REVIEWED.xlsx \
      --by "Reviewer Name" --on 2026-10-05 --version satts-ar-1.0.0 [--write] [--approve-manifest]
"""
import argparse
import json
import os

from openpyxl import load_workbook

ROOT = os.path.join(os.path.dirname(__file__), '..', '..')
PACK = os.path.join(ROOT, 'lib', 'speech', 'scientific', 'policy', 'ar-v1')


def entry_id(file, table, entry):
    roles = entry.get('roles') or table.get('roles') or []
    domains = entry.get('domains') or table.get('domains') or []
    return f"{file}|{','.join(roles)}|{','.join(domains)}|{entry['key']}"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('workbook')
    parser.add_argument('--by', required=True, help='reviewer name (approvedBy)')
    parser.add_argument('--on', required=True, help='review date YYYY-MM-DD (approvedOn)')
    parser.add_argument('--version', required=True, help='new policyVersion, e.g. satts-ar-1.0.0')
    parser.add_argument('--write', action='store_true')
    parser.add_argument('--approve-manifest', action='store_true')
    args = parser.parse_args()

    answers = {}
    wb = load_workbook(args.workbook, read_only=True)
    for ws in wb.worksheets:
        header = None
        for row in ws.iter_rows(values_only=True):
            if header is None:
                header = list(row)
                if 'المعرّف' not in header:
                    break
                continue
            record = dict(zip(header, row))
            key = record.get('المعرّف')
            if not key:
                continue
            answers[key] = (str(record.get('موافق؟') or '').strip(), str(record.get('التصحيح') or '').strip())

    manifest_path = os.path.join(PACK, 'manifest.json')
    manifest = json.load(open(manifest_path, encoding='utf-8'))
    changes, remaining = [], 0
    tables = {}
    for file in manifest['files']:
        table = json.load(open(os.path.join(PACK, file), encoding='utf-8'))
        tables[file] = table
        for entry in table['entries']:
            verdict, correction = answers.get(entry_id(file, table, entry), ('', ''))
            if verdict == 'نعم' and entry['status'] != 'approved':
                changes.append((file, entry['key'], 'approve', entry['natural']))
                entry.update(status='approved', approvedBy=args.by, approvedOn=args.on)
            elif verdict == 'لا' and correction:
                changes.append((file, entry['key'], f"correct «{entry['natural']}» →", correction))
                entry.update(natural=correction, status='approved', approvedBy=args.by, approvedOn=args.on)
            if entry['status'] != 'approved':
                remaining += 1

    for file, key, action, text in changes:
        print(f'{file}:{key}  {action} «{text}»')
    print(f'{len(changes)} change(s); {remaining} entr(y/ies) still proposed')
    if not args.write:
        print('dry run: nothing written (pass --write)')
        return
    for file, table in tables.items():
        with open(os.path.join(PACK, file), 'w', encoding='utf-8') as f:
            f.write(json.dumps(table, ensure_ascii=False, indent=2) + '\n')
    manifest['policyVersion'] = args.version
    if args.approve_manifest:
        if remaining:
            raise SystemExit(f'refusing --approve-manifest: {remaining} entries are still proposed')
        manifest['status'] = 'approved'
    with open(manifest_path, 'w', encoding='utf-8') as f:
        f.write(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n')
    print('written; now run: npx tsx scripts/satts/relock-policy.ts')


if __name__ == '__main__':
    main()
